"""
apply_theme.py

Helper functions for building on-brand .pptx decks from a theme.json config.
This is a starting point meant to be adapted per deck, not a black-box CLI.

Requires: python-pptx and Pillow  (pip install python-pptx pillow)

Typical usage from a generation script:

    from apply_theme import load_theme, new_presentation, add_title_slide, \
        add_section_slide, add_content_slide, add_chart_slide

    theme = load_theme("theme.json")
    prs = new_presentation(theme)

    add_title_slide(prs, theme, "Q3 Business Review", "Acme Corp — October 2026")
    add_section_slide(prs, theme, "Performance")
    add_content_slide(prs, theme, "Revenue grew 34% YoY, driven by EMEA",
                       bullets=["EMEA up 61%", "APAC flat", "Churn down 2pts"])
    add_chart_slide(prs, theme, "Monthly active users trending up",
                     categories=["Jan","Feb","Mar","Apr"],
                     series={"MAU": [120, 135, 150, 171]})

    prs.save("output/q3_review.pptx")
"""

import json
import os
import re
import warnings
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER


def load_theme(path):
    """Load theme.json and resolve its asset paths (template_pptx, logo.path)
    relative to the theme file's own directory — not the process's current
    working directory. Absolute paths are left untouched."""
    theme_dir = os.path.dirname(os.path.abspath(path))

    with open(path, "r", encoding="utf-8") as f:
        theme = json.load(f)

    def _resolve(rel_path):
        if not rel_path or os.path.isabs(rel_path):
            return rel_path
        return os.path.normpath(os.path.join(theme_dir, rel_path))

    if theme.get("template_pptx"):
        theme["template_pptx"] = _resolve(theme["template_pptx"])

    logo = theme.get("logo")
    if logo and logo.get("path"):
        logo["path"] = _resolve(logo["path"])

    _validate_theme(theme, path)
    return theme


_HEX_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")


def _validate_theme(theme, source_path="theme.json"):
    """Fail fast with a clear message rather than mid-render. Checks are
    deliberately non-exhaustive — this catches the mistakes that are both
    common and silent (bad hex, empty chart_sequence, missing files) without
    being a full JSON-schema validator."""
    errors = []

    colors = theme.get("colors", {})
    if not colors:
        errors.append("theme.colors is missing or empty.")
    else:
        for key in ("primary", "secondary", "accent", "background", "text_primary"):
            val = colors.get(key)
            if val is not None and not _HEX_RE.match(val):
                errors.append(f"colors.{key} = {val!r} is not a valid 6-digit hex color (e.g. '#0B3D91').")
        chart_sequence = colors.get("chart_sequence")
        if chart_sequence is not None:
            if not isinstance(chart_sequence, list) or len(chart_sequence) == 0:
                errors.append("colors.chart_sequence must be a non-empty list of hex colors.")
            else:
                for c in chart_sequence:
                    if not _HEX_RE.match(c):
                        errors.append(f"colors.chart_sequence contains an invalid hex color: {c!r}.")

    layouts = theme.get("layouts")
    if layouts is not None:
        valid = {"title", "section", "content", "chart"}
        if not isinstance(layouts, dict):
            errors.append("theme.layouts must be an object mapping slide kinds to layout names.")
        else:
            for k, v in layouts.items():
                if k not in valid or not isinstance(v, str) or not v.strip():
                    errors.append(f"layouts.{k} = {v!r}: keys must be one of {sorted(valid)} and values layout names.")

    fonts = theme.get("fonts", {})
    if not fonts.get("heading") or not fonts.get("body"):
        errors.append("fonts.heading and fonts.body should both be set (fall back to a common font like Arial if unsure).")

    # Missing template/logo files are a deliberate graceful-fallback case
    # elsewhere in this module (new_presentation / _add_logo skip them
    # silently), so flag them as a warning — loud enough to notice, but not
    # a hard failure, since "no logo yet" or "no template yet" is valid.
    template_path = theme.get("template_pptx")
    if template_path and not os.path.exists(template_path):
        warnings.warn(
            f"theme.template_pptx points to '{template_path}', which does not "
            f"exist (resolved from {source_path}) — falling back to a blank "
            f"presentation.",
            stacklevel=3,
        )

    logo = theme.get("logo")
    if logo and logo.get("path") and not os.path.exists(logo["path"]):
        warnings.warn(
            f"theme.logo.path points to '{logo['path']}', which does not "
            f"exist (resolved from {source_path}) — no logo will be rendered.",
            stacklevel=3,
        )

    if errors:
        raise ValueError(
            "Invalid theme.json (" + source_path + "):\n- " + "\n- ".join(errors)
        )


def _rgb(hex_str):
    hex_str = hex_str.lstrip("#")
    return RGBColor.from_string(hex_str.upper())


def _style(paragraph, size=None, bold=None, name=None, color=None, alignment=None):
    """Apply character formatting to the paragraph's runs (call after setting
    the text). Paragraph-level `font` only writes defRPr, which PowerPoint
    ignores for existing runs, so format the runs themselves."""
    for run in paragraph.runs:
        f = run.font
        if size is not None:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if name:
            f.name = name
        if color:
            f.color.rgb = _rgb(color)
    if alignment is not None:
        paragraph.alignment = alignment


def _set_bullet(paragraph, indent_pt=18):
    """Real bullet with a hanging indent (wrapped lines align with the text)."""
    from pptx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent_pt))))
    pPr.set("indent", str(-int(Pt(indent_pt))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


def _margins(theme):
    m = theme.get("layout_rules", {}).get("margins_in", {})
    return m.get("left", 0.6), m.get("right", 0.6), m.get("top", 0.5)


def new_presentation(theme, preserve_template_slides=False):
    """Start from the company's own template if available (preserves real
    master slides/layouts), otherwise start blank.

    preserve_template_slides: if the template_pptx already contains slides
    (sample content, a cover the user wants kept, slides that other shapes
    depend on, etc.), pass True to keep them and simply append new slides
    after them. Defaults to False (strip them) because most template files
    used with this skill are pure "layout libraries" — but this is now an
    explicit, deliberate choice rather than an unconditional wipe."""
    template_path = theme.get("template_pptx")
    if template_path and os.path.exists(template_path):
        prs = Presentation(template_path)
        if not preserve_template_slides:
            xml_slides = prs.slides._sldIdLst
            for sld_id in list(xml_slides):
                # Drop the relationship too, otherwise the slide part stays in
                # the package as an orphan (bloated file, PowerPoint repair prompt).
                prs.part.drop_rel(sld_id.rId)
                xml_slides.remove(sld_id)
        # Preserve the template's own slide geometry (e.g. 4:3 or a custom
        # size) — do NOT override it with the default 16:9 dimensions.
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    return prs


def _blank_layout(prs):
    """Find the closest thing to a 'blank' layout, rather than assuming it
    lives at index 6 — that's only true for PowerPoint's default template,
    and custom template_pptx files can order/name layouts differently.

    Note even PowerPoint's own default 'Blank' layout carries a few
    inherited (invisible until used) placeholders such as date/footer/slide
    number, so we can't require exactly zero — we prefer a layout literally
    named "blank", and otherwise fall back to whichever layout has the
    fewest placeholders."""
    named_candidates = [
        layout for layout in prs.slide_layouts
        if "blank" in layout.name.strip().lower()
    ]
    if named_candidates:
        return named_candidates[0]

    if len(prs.slide_layouts) == 0:
        raise ValueError(
            "This template_pptx has no slide layouts at all, so there is "
            "nothing to build slides from. Use a different template or "
            "fall back to a blank Presentation()."
        )

    return min(prs.slide_layouts, key=lambda layout: len(layout.placeholders))


# --- Template layout mode ---------------------------------------------------
# When theme.json has a "layouts" mapping (kind -> layout name in template_pptx),
# slides are created from those real layouts and the title/body placeholders
# are filled, so fonts, colors, logo and footer come from the template's master.
# Without a mapping (or if a name isn't found) the drawn-shape fallback is used.

_TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)


def _mapped_layout(prs, theme, kind):
    """Return the template layout configured for `kind` ("title", "section",
    "content", "chart"), or None to use the drawn fallback."""
    name = (theme.get("layouts") or {}).get(kind)
    if not name:
        return None
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == str(name).strip().lower():
            return layout
    warnings.warn(
        f"theme.layouts.{kind} = {name!r} was not found in the template "
        f"(available: {[l.name for l in prs.slide_layouts]}) — using the "
        f"drawn fallback for this slide type.",
        stacklevel=3,
    )
    return None


def _placeholder(slide, types):
    for ph in slide.placeholders:
        if ph.placeholder_format.type in types:
            return ph
    return None


def _prune_empty_placeholders(slide):
    """Remove placeholders left empty so no 'Click to add text' ghosts remain."""
    for ph in list(slide.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)


def _set_title(slide, text):
    ph = _placeholder(slide, _TITLE_TYPES)
    if ph is not None:
        ph.text_frame.text = text
    else:
        warnings.warn("Layout has no title placeholder; the title was not rendered.", stacklevel=3)


def _layout_footer(slide, prs, theme, page_number):
    """Footer text / page number from theme rules; logo comes from the master."""
    _add_footer_and_logo(slide, prs, theme, page_number=page_number, logo=False)


def _logo_dimensions(logo_path, max_w):
    """Return (width, height) EMU for the logo, preserving its real aspect
    ratio so non-square logos (most logos) don't get stretched. Falls back
    to a square box if the image's intrinsic size can't be read."""
    try:
        from PIL import Image
        with Image.open(logo_path) as img:
            px_w, px_h = img.size
        if px_w > 0 and px_h > 0:
            height = Emu(int(max_w * px_h / px_w))
            return max_w, height
    except Exception:
        pass
    return max_w, max_w


def _add_logo(slide, prs, theme):
    """Place the logo per theme.logo.placement, sized from its real aspect
    ratio. Shared by every slide type so placement is consistent. Bottom
    placements sit above the footer/page-number row so they never overlap it."""
    logo = theme.get("logo", {})
    logo_path = logo.get("path")
    if not (logo_path and os.path.exists(logo_path)):
        return
    max_w = Inches(logo.get("max_width_in", 0.9))
    logo_w, logo_h = _logo_dimensions(logo_path, max_w)
    placement = logo.get("placement", "bottom-right")
    if placement == "bottom-right":
        left = prs.slide_width - logo_w - Inches(0.3)
        top = prs.slide_height - logo_h - Inches(0.5)
    elif placement == "top-right":
        left = prs.slide_width - logo_w - Inches(0.3)
        top = Inches(0.2)
    elif placement == "bottom-left":
        left, top = Inches(0.3), prs.slide_height - logo_h - Inches(0.5)
    else:  # "top-left" and any unrecognized value fall back here
        left, top = Inches(0.3), Inches(0.2)
    slide.shapes.add_picture(logo_path, left, top, width=logo_w, height=logo_h)


def _add_footer_and_logo(slide, prs, theme, page_number=None, logo=True):
    rules = theme.get("layout_rules", {})
    colors = theme.get("colors", {})

    if rules.get("footer_text"):
        left, top, width, height = Inches(0.6), prs.slide_height - Inches(0.4), Inches(6), Inches(0.3)
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.text = rules["footer_text"]
        _style(tf.paragraphs[0], size=theme.get("type_scale", {}).get("footnote_pt", 10),
               name=theme.get("fonts", {}).get("body", "Arial"),
               color=colors.get("text_primary", "#1A1A1A"))

    if rules.get("show_page_numbers"):
        # Honor the theme's show_page_numbers setting even if the caller
        # didn't pass an explicit page_number — derive it from the slide's
        # actual position in the deck so far.
        resolved_number = page_number
        if resolved_number is None:
            resolved_number = len(prs.slides)
        left = prs.slide_width - Inches(1)
        top = prs.slide_height - Inches(0.4)
        box = slide.shapes.add_textbox(left, top, Inches(0.6), Inches(0.3))
        tf = box.text_frame
        tf.text = str(resolved_number)
        _style(tf.paragraphs[0], size=theme.get("type_scale", {}).get("footnote_pt", 10),
               name=theme.get("fonts", {}).get("body", "Arial"),
               color=colors.get("text_primary", "#1A1A1A"), alignment=PP_ALIGN.RIGHT)

    if logo:
        _add_logo(slide, prs, theme)


def add_title_slide(prs, theme, title, subtitle=None):
    colors = theme.get("colors", {})
    fonts = theme.get("fonts", {})
    type_scale = theme.get("type_scale", {})

    layout = _mapped_layout(prs, theme, "title")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        _set_title(slide, title)
        sub_ph = _placeholder(slide, (PP_PLACEHOLDER.SUBTITLE,) + _BODY_TYPES)
        if subtitle and sub_ph is not None:
            sub_ph.text_frame.text = subtitle
        _prune_empty_placeholders(slide)
        return slide

    slide = prs.slides.add_slide(_blank_layout(prs))
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(colors.get("primary", "#0B3D91"))

    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), prs.slide_width - Inches(1.6), Inches(1.5))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = title
    _style(tf.paragraphs[0], size=type_scale.get("title_slide_pt", 40), bold=True,
           name=fonts.get("heading", "Arial"), color=colors.get("text_on_dark", "#FFFFFF"))

    if subtitle:
        sbox = slide.shapes.add_textbox(Inches(0.8), Inches(4.1), prs.slide_width - Inches(1.6), Inches(0.8))
        stf = sbox.text_frame
        stf.word_wrap = True
        stf.text = subtitle
        _style(stf.paragraphs[0], size=18, name=fonts.get("body", "Arial"),
               color=colors.get("text_on_dark", "#FFFFFF"))

    _add_logo(slide, prs, theme)

    return slide


def add_section_slide(prs, theme, section_title):
    colors = theme.get("colors", {})
    fonts = theme.get("fonts", {})

    layout = _mapped_layout(prs, theme, "section")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        _set_title(slide, section_title)
        _prune_empty_placeholders(slide)
        return slide

    slide = prs.slides.add_slide(_blank_layout(prs))
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(colors.get("secondary", colors.get("primary", "#0B3D91")))

    box = slide.shapes.add_textbox(Inches(0.8), prs.slide_height // 2 - Inches(0.6), prs.slide_width - Inches(1.6), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = section_title
    _style(tf.paragraphs[0], size=32, bold=True, name=fonts.get("heading", "Arial"),
           color=colors.get("text_on_dark", "#FFFFFF"))
    _add_logo(slide, prs, theme)
    return slide


def _add_slide_title(slide, prs, theme, title):
    """Slide title box shared by content and chart slides (honors margins)."""
    colors = theme.get("colors", {})
    fonts = theme.get("fonts", {})
    m_left, m_right, m_top = _margins(theme)
    tbox = slide.shapes.add_textbox(Inches(m_left), Inches(m_top),
                                     prs.slide_width - Inches(m_left + m_right), Inches(1))
    ttf = tbox.text_frame
    ttf.word_wrap = True
    ttf.text = title
    _style(ttf.paragraphs[0], size=theme.get("type_scale", {}).get("slide_title_pt", 28),
           bold=True, name=fonts.get("heading", "Arial"), color=colors.get("primary", "#0B3D91"))


def add_content_slide(prs, theme, title, bullets=None, page_number=None):
    colors = theme.get("colors", {})
    fonts = theme.get("fonts", {})
    type_scale = theme.get("type_scale", {})
    m_left, m_right, m_top = _margins(theme)

    if bullets and len(bullets) > 6:
        warnings.warn(
            f"add_content_slide('{title}'): {len(bullets)} bullets given, "
            f"only the first 6 will be rendered (the 6x6 guideline is a "
            f"content ceiling, not a silent-truncation rule — split this "
            f"into two slides or trim the content instead).",
            stacklevel=2,
        )

    layout = _mapped_layout(prs, theme, "content")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        _set_title(slide, title)
        body = _placeholder(slide, _BODY_TYPES)
        if bullets and body is not None:
            tf = body.text_frame
            for i, bullet in enumerate(bullets[:6]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = bullet  # bullet style/fonts inherited from the layout
        elif bullets:
            warnings.warn("Layout has no body placeholder; bullets were not rendered.", stacklevel=2)
        _prune_empty_placeholders(slide)
        _layout_footer(slide, prs, theme, page_number)
        return slide

    slide = prs.slides.add_slide(_blank_layout(prs))
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(colors.get("background", "#FFFFFF"))

    _add_slide_title(slide, prs, theme, title)

    if bullets:
        bbox = slide.shapes.add_textbox(Inches(m_left), Inches(1.7),
                                         prs.slide_width - Inches(m_left + m_right), Inches(4.5))
        btf = bbox.text_frame
        btf.word_wrap = True
        for i, bullet in enumerate(bullets[:6]):  # enforce 6x6-ish ceiling
            p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
            p.text = bullet
            _style(p, size=type_scale.get("body_pt", 16), name=fonts.get("body", "Arial"),
                   color=colors.get("text_primary", "#1A1A1A"))
            _set_bullet(p)
            p.space_after = Pt(10)

    _add_footer_and_logo(slide, prs, theme, page_number=page_number)
    return slide


def add_chart_slide(prs, theme, title, categories, series, chart_type="bar", page_number=None):
    """series: dict of {series_name: [values]} aligned with categories."""
    colors = theme.get("colors", {})
    fonts = theme.get("fonts", {})
    type_scale = theme.get("type_scale", {})

    m_left, m_right, _ = _margins(theme)
    x, y, cx, cy = Inches(m_left), Inches(1.7), prs.slide_width - Inches(m_left + m_right), Inches(4.8)

    layout = _mapped_layout(prs, theme, "chart")
    if layout is not None:
        slide = prs.slides.add_slide(layout)
        _set_title(slide, title)
        body = _placeholder(slide, _BODY_TYPES)
        if body is not None:  # use the layout's content area for the chart
            x, y, cx, cy = body.left, body.top, body.width, body.height
            body._element.getparent().remove(body._element)
        _prune_empty_placeholders(slide)
    else:
        slide = prs.slides.add_slide(_blank_layout(prs))
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(colors.get("background", "#FFFFFF"))
        _add_slide_title(slide, prs, theme, title)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series.items():
        chart_data.add_series(name, values)

    xl_type = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
    }.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)

    gframe = slide.shapes.add_chart(xl_type, x, y, cx, cy, chart_data)
    chart = gframe.chart
    chart.font.name = fonts.get("body", "Arial")
    chart.font.size = Pt(type_scale.get("footnote_pt", 10) + 2)
    chart.font.color.rgb = _rgb(colors.get("text_primary", "#1A1A1A"))
    chart.has_legend = len(series) > 1 or chart_type == "pie"
    if chart.has_legend:
        from pptx.enum.chart import XL_LEGEND_POSITION
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if chart_type == "pie" and len(series) > 1:
        warnings.warn("add_chart_slide: pie charts render only the first series.", stacklevel=2)

    chart_sequence = colors.get("chart_sequence", [colors.get("primary", "#0B3D91")])

    if chart_type == "line":
        # A series fill doesn't affect a line's stroke — set the line color
        # (and marker fill, so markers match) per series instead.
        for i, plot_series in enumerate(chart.plots[0].series):
            color = chart_sequence[i % len(chart_sequence)]
            plot_series.format.line.color.rgb = _rgb(color)
            plot_series.format.line.width = Pt(2.25)
            try:
                plot_series.marker.format.fill.solid()
                plot_series.marker.format.fill.fore_color.rgb = _rgb(color)
            except (AttributeError, TypeError):
                pass  # not all chart engines expose marker formatting the same way
    elif chart_type == "pie":
        # A series-level fill would make every slice the same color — color
        # each point (slice) individually from chart_sequence instead.
        plot_series = chart.plots[0].series[0]
        for i, point in enumerate(plot_series.points):
            color = chart_sequence[i % len(chart_sequence)]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(color)
    else:
        for i, plot_series in enumerate(chart.plots[0].series):
            color = chart_sequence[i % len(chart_sequence)]
            plot_series.format.fill.solid()
            plot_series.format.fill.fore_color.rgb = _rgb(color)

    _add_footer_and_logo(slide, prs, theme, page_number=page_number, logo=layout is None)
    return slide
